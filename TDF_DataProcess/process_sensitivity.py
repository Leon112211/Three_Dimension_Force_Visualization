"""
Sensor Sensitivity Analysis
- Reads raw CSVs (x1, y1, z1 field readings in uT at each force step)
- Computes per-file means, rewrites the Computed_Means summary sheet
- Calculates sensitivity (uT/N) via least-squares regression, with R2
- Plots main + coupling axes per sensor per applied-force axis
- Writes coupling sensitivities to Processed_Data.xlsx
"""

import time
import shutil
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

matplotlib.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei', 'DejaVu Sans']
matplotlib.rcParams['axes.unicode_minus'] = False

BASE = Path(__file__).resolve().parent   # repo root = folder containing this script
RAW  = BASE / 'Raw_Data'

SUMMARY_PATH = RAW / 'xyz_H2H4H6数据汇总.xlsx'
MEANS_SHEET  = 'Computed_Means'
PD_PATH      = BASE / 'Processed_Data.xlsx'
PLOTS_DIR    = BASE / 'Plots'

R2_WARN_THRESHOLD = 0.98   # main-channel fits below this get a linearity warning

# ── File structure config ───────────────────────────────────────────────────
AXES_CONFIG = {
    'X': {
        'dir':     RAW / '2025_11_11 x轴H2H4H6数据/11.11_x轴',
        'prefix':  '11.11_x轴',
        'sensors': ['H2', 'H4', 'H6'],
        'forces':  [i * 0.5 for i in range(7)],   # 0 – 3 N
    },
    'Y': {
        'dir':     RAW / '2025_11_10 y轴H2H6数据/11.10_y轴',
        'prefix':  '11.10_y轴',
        'sensors': ['H2', 'H6'],                   # H4 not measured
        'forces':  [i * 0.5 for i in range(7)],   # 0 – 3 N
    },
    'Z': {
        'dir':     RAW / '2025_11_12 z轴H2H4H6/11.12_Z',
        'prefix':  '11.12_Z',
        'sensors': ['H2', 'H4', 'H6'],
        'forces':  [i * 0.5 for i in range(41)],  # 0 – 20 N
    },
}


# ── Helpers ─────────────────────────────────────────────────────────────────
def backup_before_save(path: Path) -> None:
    """Keep a rolling .bak copy of an output file before overwriting it."""
    if path.exists():
        shutil.copy2(path, path.with_name(path.name + '.bak'))


def force_to_filename(f: float) -> str:
    """0.5 → '0_5', 1.0 → '1', 1.5 → '1_5', 0.25 → '0_25', 20.0 → '20'"""
    if f == int(f):
        return str(int(f))
    return f"{f:g}".replace('.', '_')


def read_raw_csv(fpath: Path) -> pd.DataFrame:
    """Load one raw capture (columns x1,y1,z1 in uT), dropping the summary
    row the logger appends at the end of every file.

    That last row holds column MEANS, not a real sample. In most files only
    the loaded axis is filled (other cells empty -> NaN, removed by dropna);
    in the remaining files it is complete and equals the mean of all rows
    above it, written with ~10 significant digits while real samples have
    exactly 2 decimals - both properties are used to detect and drop it.
    """
    df = pd.read_csv(fpath).dropna().reset_index(drop=True)
    if len(df) >= 3:
        body, last = df.iloc[:-1], df.iloc[-1]
        looks_like_mean = all(abs(last[c] - body[c].mean()) <= 1e-3 for c in df.columns)
        off_sample_grid = any(abs(last[c] * 100 - round(last[c] * 100)) > 1e-6
                              for c in df.columns)
        if looks_like_mean and off_sample_grid:
            df = body.reset_index(drop=True)
    return df


def least_squares_fit(forces, readings):
    """OLS fit y = a*x + b on level means.
    Returns dict(slope, intercept, r2, max_resid): slope in uT/N,
    max_resid = largest |fit residual| across force levels, in uT."""
    x = np.array(forces, dtype=float)
    y = np.array(readings, dtype=float)
    A = np.column_stack([x, np.ones_like(x)])
    slope, intercept = np.linalg.lstsq(A, y, rcond=None)[0]
    resid = y - (slope * x + intercept)
    ss_tot = float(np.sum((y - y.mean()) ** 2))
    r2 = 1.0 - float(np.sum(resid ** 2)) / ss_tot if ss_tot > 0 else float('nan')
    return {'slope': float(slope), 'intercept': float(intercept),
            'r2': r2, 'max_resid': float(np.abs(resid).max())}


# ── Step 1: Collect all per-file means ──────────────────────────────────────
def collect_means() -> pd.DataFrame:
    print("【步骤 1/5】读取原始数据")
    records = []
    missing = []
    for axis, cfg in AXES_CONFIG.items():
        n_before = len(records)
        for sensor in cfg['sensors']:
            sensor_dir = cfg['dir'] / f"{cfg['prefix']}_{sensor}"
            for force in cfg['forces']:
                fpath = sensor_dir / f"{force_to_filename(force)}.csv"
                if not fpath.exists():
                    missing.append(fpath)
                    print(f"  缺失文件: {fpath}")
                    continue
                df = read_raw_csv(fpath)
                records.append({
                    'Sensor':       sensor,
                    'Applied_Axis': axis,
                    'Force_N':      force,
                    'Mean_X':       df['x1'].mean(),
                    'Mean_Y':       df['y1'].mean(),
                    'Mean_Z':       df['z1'].mean(),
                })
        got  = len(records) - n_before
        exp  = len(cfg['sensors']) * len(cfg['forces'])
        fmax = max(cfg['forces'])
        print(f"  {axis} 轴: {' '.join(cfg['sensors']):<8} x {len(cfg['forces']):>2} 力级"
              f" (0-{fmax:g} N)  -> {got}/{exp} 个文件")

    means_df = pd.DataFrame(records)
    expected = sum(len(cfg['sensors']) * len(cfg['forces']) for cfg in AXES_CONFIG.values())
    print(f"  合计: {len(means_df)} / {expected} 个文件读取成功")
    if missing:
        raise SystemExit(
            f"错误: 缺失 {len(missing)} 个原始 CSV 文件 - 拒绝用不完整的数据计算灵敏度,"
            f"请检查上方的'缺失文件'列表。")
    return means_df


# ── Step 2: Rewrite the Computed_Means sheet in the summary xlsx ────────────
def write_means_sheet(means_df: pd.DataFrame) -> None:
    print("\n【步骤 2/5】写入汇总表")
    wb = load_workbook(SUMMARY_PATH)

    # Idempotent: regenerate a dedicated sheet from scratch on every run instead
    # of appending blocks to the original data sheet (which used to grow forever).
    if MEANS_SHEET in wb.sheetnames:
        del wb[MEANS_SHEET]
    ws = wb.create_sheet(MEANS_SHEET)

    header_fill = PatternFill('solid', fgColor='D9E1F2')
    header_font = Font(bold=True)

    ws.cell(1, 1, 'Computed per-step means - regenerated by process_sensitivity.py on every run').font = header_font
    start_row = 3

    for axis in ['X', 'Y', 'Z']:
        block = means_df[means_df['Applied_Axis'] == axis]
        sensors_in_block = sorted(block['Sensor'].unique())

        # Block title row
        ws.cell(start_row, 1, f'Applied axis: {axis}').font = Font(bold=True)
        start_row += 1

        # Column headers: Force_N | sensor_Mean_X | sensor_Mean_Y | sensor_Mean_Z | ...
        col = 1
        hcell = ws.cell(start_row, col, 'Force (N)')
        hcell.font = header_font
        hcell.fill = header_fill
        col += 1
        for sensor in sensors_in_block:
            for reading in ['Mean_X', 'Mean_Y', 'Mean_Z']:
                hcell = ws.cell(start_row, col, f'{sensor}_{reading}')
                hcell.font = header_font
                hcell.fill = header_fill
                col += 1
        start_row += 1

        # Data rows (pivot by force)
        forces = sorted(block['Force_N'].unique())
        for force in forces:
            col = 1
            ws.cell(start_row, col, force)
            col += 1
            for sensor in sensors_in_block:
                row_data = block[(block['Force_N'] == force) & (block['Sensor'] == sensor)]
                if row_data.empty:
                    col += 3
                    continue
                ws.cell(start_row, col,   round(row_data.iloc[0]['Mean_X'], 4))
                ws.cell(start_row, col+1, round(row_data.iloc[0]['Mean_Y'], 4))
                ws.cell(start_row, col+2, round(row_data.iloc[0]['Mean_Z'], 4))
                col += 3
            start_row += 1

        start_row += 1  # blank between axis blocks

    backup_before_save(SUMMARY_PATH)
    wb.save(SUMMARY_PATH)
    print(f"  已重写工作表 '{MEANS_SHEET}' -> {SUMMARY_PATH.name}")


# ── Step 3: Least-squares sensitivity calculation ───────────────────────────
def compute_sensitivities(means_df: pd.DataFrame):
    """Returns (sensitivity, fits):
    sensitivity[sensor][applied_axis][reading_axis] = slope (uT/N)
    fits[sensor][applied_axis][reading_axis]        = full fit dict"""
    print("\n【步骤 3/5】最小二乘灵敏度拟合 (uT/N)")
    print("  传感器 施力轴      X 通道      Y 通道      Z 通道   主通道R2  最大残差")
    sensitivity = {}
    fits = {}
    any_warn = False

    for sensor in ['H2', 'H4', 'H6']:
        sensitivity[sensor] = {}
        fits[sensor] = {}
        for axis, cfg in AXES_CONFIG.items():
            if sensor not in cfg['sensors']:
                continue
            block = means_df[(means_df['Sensor'] == sensor) & (means_df['Applied_Axis'] == axis)]
            block = block.sort_values('Force_N')
            forces = block['Force_N'].values
            slopes, quality = {}, {}
            for reading_ax, col in [('X', 'Mean_X'), ('Y', 'Mean_Y'), ('Z', 'Mean_Z')]:
                fit = least_squares_fit(forces, block[col].values)
                slopes[reading_ax]  = fit['slope']
                quality[reading_ax] = fit
            sensitivity[sensor][axis] = slopes
            fits[sensor][axis] = quality

            main = quality[axis]
            resid_N = main['max_resid'] / abs(main['slope']) if main['slope'] else float('inf')
            warn = main['r2'] < R2_WARN_THRESHOLD
            any_warn = any_warn or warn
            print(f"  {sensor:<5}  {axis:^4}  {slopes['X']:>+10.3f}  {slopes['Y']:>+10.3f}"
                  f"  {slopes['Z']:>+10.3f}    {main['r2']:.4f}   {resid_N:>5.2f} N"
                  f"{'  [!]' if warn else ''}")

    if any_warn:
        print(f"  [!] = 主通道 R2 < {R2_WARN_THRESHOLD}, 线性模型在该量程内可能不成立")

    return sensitivity, fits


# ── Step 4: Visualization ───────────────────────────────────────────────────
AXIS_COLOR   = {'X': '#D62728', 'Y': '#1F77B4', 'Z': '#2CA02C'}
AXIS_MARKER  = {'X': 'o',       'Y': 's',       'Z': '^'}
SENSORS      = ['H2', 'H4', 'H6']


def create_plots(means_df: pd.DataFrame, sensitivity: dict, fits: dict) -> None:
    print("\n【步骤 4/5】生成图表")
    PLOTS_DIR.mkdir(exist_ok=True)
    plt.style.use('seaborn-v0_8-whitegrid')

    for sensor in SENSORS:
        fig, axes_plot = plt.subplots(1, 3, figsize=(17, 5.2))
        fig.patch.set_facecolor('#FAFAFA')

        fig.suptitle(
            f'Sensor {sensor}  |  Axis Sensitivity Analysis (Least-Squares Regression)',
            fontsize=13, fontweight='bold', y=1.01, color='#1a1a1a'
        )

        for ax_idx, applied_ax in enumerate(['X', 'Y', 'Z']):
            ax = axes_plot[ax_idx]
            ax.set_facecolor('#F7F9FC')

            # ── no data panel ──
            if sensor not in AXES_CONFIG[applied_ax]['sensors']:
                for spine in ax.spines.values():
                    spine.set_edgecolor('#cccccc')
                ax.text(0.5, 0.5, 'Not measured',
                        ha='center', va='center', transform=ax.transAxes,
                        fontsize=11, color='#aaaaaa', style='italic')
                ax.set_title(f'Applied force: {applied_ax}-axis', fontsize=10,
                             fontweight='bold', color='#555555')
                ax.set_xticks([]); ax.set_yticks([])
                continue

            block = means_df[
                (means_df['Sensor'] == sensor) &
                (means_df['Applied_Axis'] == applied_ax)
            ].sort_values('Force_N')

            forces        = block['Force_N'].values
            coupling_axes = [a for a in ['X', 'Y', 'Z'] if a != applied_ax]

            for reading_ax in [applied_ax] + coupling_axes:
                col    = f'Mean_{reading_ax}'
                values = block[col].values
                slope  = sensitivity[sensor][applied_ax][reading_ax]
                is_main = (reading_ax == applied_ax)

                color  = AXIS_COLOR[reading_ax]
                marker = AXIS_MARKER[reading_ax]
                lw     = 2.2 if is_main else 1.4
                ls     = '-'  if is_main else '--'
                role   = 'Primary' if is_main else 'Coupling'
                label  = f'{reading_ax.lower()}1  [{role}]   k = {slope:+.2f} uT/N'
                if is_main:
                    label += f'   R2 = {fits[sensor][applied_ax][reading_ax]["r2"]:.4f}'
                alpha_scatter = 0.50 if is_main else 0.35
                ms     = 36   if is_main else 22

                # raw mean scatter
                ax.scatter(forces, values,
                           color=color, marker=marker,
                           s=ms, alpha=alpha_scatter, zorder=4, linewidths=0.5,
                           edgecolors=color)

                # least-squares fitted line (reuse the Step 3 fit)
                fit    = fits[sensor][applied_ax][reading_ax]
                f_fit  = np.linspace(forces.min(), forces.max(), 300)
                ax.plot(f_fit, fit['slope'] * f_fit + fit['intercept'],
                        color=color, lw=lw, ls=ls, label=label, zorder=5)

            # ── axis decorations ──
            coupling_str = ' & '.join(f'{a}-axis' for a in coupling_axes)
            ax.set_title(
                f'Applied: {applied_ax}-axis   |   Coupling: {coupling_str}',
                fontsize=9.5, fontweight='bold', color='#333333', pad=8
            )
            ax.set_xlabel('Applied Force (N)', fontsize=9, labelpad=5)
            ax.set_ylabel('Sensor Output (uT)', fontsize=9, labelpad=5)
            ax.tick_params(labelsize=8)

            legend = ax.legend(
                fontsize=8, loc='best',
                framealpha=0.92, edgecolor='#cccccc',
                handlelength=2.4, labelspacing=0.5
            )
            for text in legend.get_texts():
                text.set_color('#222222')

            ax.grid(True, alpha=0.4, linewidth=0.6)
            for spine in ax.spines.values():
                spine.set_edgecolor('#cccccc')
                spine.set_linewidth(0.8)

        plt.tight_layout(rect=[0, 0, 1, 0.98])
        save_path = PLOTS_DIR / f'sensitivity_{sensor}.png'
        plt.savefig(save_path, dpi=160, bbox_inches='tight', facecolor=fig.get_facecolor())
        plt.close()
        print(f"  已保存: Plots/{save_path.name}")


# ── Step 5: Update Processed_Data.xlsx ──────────────────────────────────────
# Layout (1-indexed in openpyxl):
# Row 1: H2 / H4 / H6 header
# Row 2: column sub-headers
# Row 3: X applied · Row 4: Y applied · Row 5: Z applied
#
# Per sensor block = 6 columns:
#   col_offset+1: force axis label
#   col_offset+2: main sensitivity      ← fill this
#   col_offset+3: coupling axis 1 label
#   col_offset+4: coupling 1 sensitivity ← fill this
#   col_offset+5: coupling axis 2 label
#   col_offset+6: coupling 2 sensitivity ← fill this
SENSOR_COL_OFFSETS = {'H2': 0, 'H4': 6, 'H6': 12}
APPLIED_AXIS_ROWS  = {'X': 3, 'Y': 4, 'Z': 5}

# Coupling order per applied axis (as laid out in the existing file)
LAYOUT = {
    'X': {'main': 'X', 'c1': 'Y', 'c2': 'Z'},
    'Y': {'main': 'Y', 'c1': 'Z', 'c2': 'X'},
    'Z': {'main': 'Z', 'c1': 'X', 'c2': 'Y'},
}


def update_processed_data(sensitivity: dict) -> None:
    print("\n【步骤 5/5】更新 Processed_Data.xlsx")
    wb_pd = load_workbook(PD_PATH)
    ws_pd = wb_pd.active
    val_font = Font(bold=False)

    for sensor in ['H2', 'H4', 'H6']:
        col_off = SENSOR_COL_OFFSETS[sensor]
        written, na = [], []
        for applied_ax, row in APPLIED_AXIS_ROWS.items():
            if sensor not in AXES_CONFIG[applied_ax]['sensors']:
                # No data available (e.g. H4 for Y axis)
                ws_pd.cell(row, col_off + 2, 'N/A')
                ws_pd.cell(row, col_off + 4, 'N/A')
                ws_pd.cell(row, col_off + 6, 'N/A')
                na.append(applied_ax)
                continue

            layout = LAYOUT[applied_ax]
            sens = sensitivity[sensor][applied_ax]

            main_slope = round(sens[layout['main']], 4)
            c1_slope   = round(sens[layout['c1']],   4)
            c2_slope   = round(sens[layout['c2']],   4)

            # Write sensitivity values (even columns: 2, 4, 6 within the block)
            ws_pd.cell(row, col_off + 2, main_slope).font = val_font
            ws_pd.cell(row, col_off + 4, c1_slope).font   = val_font
            ws_pd.cell(row, col_off + 6, c2_slope).font   = val_font
            written.append(applied_ax)

        line = f"  {sensor}: 已写入 {'/'.join(written)} 轴"
        if na:
            line += f"  ({'/'.join(na)} 轴无数据 -> N/A)"
        print(line)

    backup_before_save(PD_PATH)
    wb_pd.save(PD_PATH)
    print(f"  已保存: {PD_PATH.name}")


def main() -> None:
    t0 = time.perf_counter()
    print("=" * 62)
    print("  三维力传感器 灵敏度标定管线  (process_sensitivity.py)")
    print("=" * 62)
    print()

    means_df = collect_means()
    write_means_sheet(means_df)
    sensitivity, fits = compute_sensitivities(means_df)
    create_plots(means_df, sensitivity, fits)
    update_processed_data(sensitivity)

    print()
    print("=" * 62)
    print(f"  完成, 耗时 {time.perf_counter() - t0:.1f} s")
    print(f"  图表目录:   {PLOTS_DIR}")
    print(f"  汇总表:     {SUMMARY_PATH}")
    print(f"  灵敏度矩阵: {PD_PATH}")
    print("=" * 62)


if __name__ == '__main__':
    main()
